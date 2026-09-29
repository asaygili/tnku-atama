"""
Web of Science "atıf yapan yayınlar" dışa aktarımından atıf klasörleri oluşturma.

WoS'ta (Citation Report → Citing articles → isteğe bağlı "Web of Science Index"
süzgeci → Export → "Tab delimited file", içerik "Full Record and Cited References")
alınan dosya okunur. Her atıf yapan yayın için:
  • kaynakçasındaki (CR) DOI'lerden – yoksa yazar/yıl/cilt/sayfadan – hangi
    yayınımıza atıf yaptığı bulunur (birden çok yayınımıza atıf → her birine),
  • WoS indeksi (WE alanı) → EK-2 kodu (SCI-E / SSCI / AHCI → 5.1, ESCI → 5.2),
  • öz atıflar (yazarlar arasında aday) ve arşivde zaten bulunan atıflar atlanır,
  • yayının atiflar\\ klasöründe ayrı bir alt klasör açılır:
        WoS_<yıl>_<ilk yazar>_<başlık> (SCI-E)\\
            kayit.json          WoS kaydı (makine için)
            endeks_bilgisi.pdf  WoS kaydından üretilen endeks bilgisi sayfası
            atif_yapan.pdf      açık erişimli tam metin (varsa)
            INDIRILECEK.txt     ücretli yayınlar için DOI bağlantısı
"""

from __future__ import annotations

import csv
import json
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from .kayit import KanitArsivi, KanitKaydi
from .ortak import doi_bul, norm, pdf_metin, slug

KAYNAK = "wos"
INDIRILECEK_LISTESI = "_indirilecek_atiflar.xlsx"

# WoS tab-delimited alan etiketleri ↔ Excel sütun adları
ALANLAR = {
    "UT": ("UT", "UT (Unique WOS ID)"), "TI": ("TI", "Article Title"),
    "AU": ("AU", "Authors"), "AF": ("AF", "Author Full Names"), "SO": ("SO", "Source Title"),
    "SN": ("SN", "ISSN"), "EI": ("EI", "eISSN"), "PY": ("PY", "Publication Year"),
    "PD": ("PD", "Publication Date"), "EA": ("EA", "Early Access Date"), "DI": ("DI", "DOI"),
    "WE": ("WE", "Web of Science Index"), "CR": ("CR", "Cited References"),
    "DT": ("DT", "Document Type"), "VL": ("VL", "Volume"), "BP": ("BP", "Start Page"),
    "AR": ("AR", "Article Number"), "PT": ("PT", "Publication Type"),
}

AYLAR = {"jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6, "jul": 7, "aug": 8,
         "sep": 9, "oct": 10, "nov": 11, "dec": 12,
         "spr": 3, "sum": 6, "fal": 9, "aut": 9, "win": 1}     # mevsim → ilk ayı


def ay_tarihi(metin: str, yil: int | None = None) -> date | None:
    """WoS tarih alanı ("MAR 2023", "MAR 15 2023", "JAN-FEB", PD + PY) → ayın ilk günü."""
    m = re.match(r"\s*([A-Za-z]{3})", metin or "")
    if not m or m.group(1).lower() not in AYLAR:
        return None
    y = re.search(r"(19|20)\d{2}", metin)
    yil = int(y.group(0)) if y else yil
    return date(yil, AYLAR[m.group(1).lower()], 1) if yil else None


ENDEKS_KODU = [("5.1", ("science citation index expanded", "sci expanded", "science citation index",
                        "social sciences citation index", "ssci", "arts humanities citation index",
                        "a hci", "ahci")),
               ("5.2", ("emerging sources citation index", "esci")),
               ("5.7", ("book citation index",))]         # BKCI kitabında atıf
ENDEKS_KISA = {"science citation index expanded": "SCI-E", "social sciences citation index": "SSCI",
               "arts humanities citation index": "AHCI", "emerging sources citation index": "ESCI",
               "book citation index": "BKCI"}


@dataclass
class WosKaydi:
    ut: str
    baslik: str
    yazarlar: list[str]
    dergi: str
    issn: str
    eissn: str
    yil: int | None
    doi: str
    endeks: str                       # WE alanı (ham)
    kaynakca: list[str] = field(default_factory=list)
    tur: str = ""
    tarih: date | None = None         # ay hassasiyetinde en erken yayım tarihi (EA / PD)
    yayin_turu: str = ""              # PT: J dergi, B kitap, S seri, P patent

    @property
    def kitap_bolumu(self) -> bool:
        return "book chapter" in self.tur.lower()

    @property
    def endeks_kodu(self) -> str | None:
        m = norm(self.endeks)
        for kod, ifadeler in ENDEKS_KODU:
            if any(norm(i) in m for i in ifadeler):
                return kod
        return None

    @property
    def endeks_kisa(self) -> str:
        m = norm(self.endeks)
        kisa = [v for k, v in ENDEKS_KISA.items() if k in m]
        return "/".join(kisa) or (self.endeks[:20] or "?")


# ── Okuma ──────────────────────────────────────────────────────────────────
def _satirlar(yol: Path) -> list[dict]:
    yol = Path(yol)
    if yol.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        ws = load_workbook(yol, read_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        basliklar = [str(h or "").strip() for h in rows[0]]
        return [{b: ("" if v is None else str(v)) for b, v in zip(basliklar, r)} for r in rows[1:]]
    if yol.suffix.lower() == ".xls":
        raise ValueError("Eski Excel (.xls) biçimi desteklenmiyor: WoS'tan 'Tab delimited file' "
                         "ya da .xlsx olarak dışa aktarın.")
    ham = yol.read_bytes()
    for kodlama in ("utf-8-sig", "utf-16", "cp1254"):
        try:
            metin = ham.decode(kodlama)
            if "\t" in metin.split("\n", 1)[0]:
                break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("Dosya okunamadı (WoS 'Tab delimited file' bekleniyor)")
    okuyucu = csv.DictReader(metin.splitlines(), delimiter="\t", quoting=csv.QUOTE_NONE)
    return [{k.strip(): (v or "") for k, v in r.items() if k} for r in okuyucu]


def oku(yol: Path) -> list[WosKaydi]:
    sonuc = []
    for r in _satirlar(yol):
        def al(etiket):
            return next((r[a].strip() for a in ALANLAR[etiket] if r.get(a)), "")
        if not (al("TI") or al("UT")):
            continue
        yil = re.search(r"(19|20)\d{2}", al("PY") or al("EA"))
        sonuc.append(WosKaydi(
            ut=al("UT"), baslik=re.sub(r"\s+", " ", al("TI")),
            yazarlar=[a.strip() for a in (al("AU") or al("AF")).split(";") if a.strip()],
            dergi=al("SO"), issn=al("SN"), eissn=al("EI"),
            yil=int(yil.group(0)) if yil else None, doi=al("DI").lower(),
            endeks=al("WE"), kaynakca=[c.strip() for c in al("CR").split(";") if c.strip()],
            tur=al("DT"), yayin_turu=al("PT"),
            tarih=min((t for t in (ay_tarihi(al("EA")),
                                   ay_tarihi(al("PD"), int(yil.group(0)) if yil else None))
                       if t), default=None)))
    return sonuc


# ── Eşleştirme ─────────────────────────────────────────────────────────────
def _ref_ipuclari(kunye: str) -> tuple[str, str, str]:
    """Künyeden (yıl, cilt, ilk sayfa)."""
    yil = (re.findall(r"\b((?:19|20)\d{2})\b", kunye) or [""])[-1]
    cilt = (re.search(r"vol\.\s*(\d+)", kunye) or [None, ""])[1]
    sayfa = (re.search(r"pp\.\s*(\d+)", kunye) or [None, ""])[1]
    return yil, cilt, sayfa


def _kisaltma_eslesir(kisaltma: str, kunye: str) -> bool:
    """WoS kaynak kısaltmasının her kelimesi künyedeki bir kelimenin başı mı?
    ("SIG PROCESS COMMUN" ↔ "Signal Processing and Communications Applications")"""
    kel = norm(kunye).split()
    parcalar = [p for p in norm(kisaltma).split() if len(p) > 1]
    return bool(parcalar) and all(any(w.startswith(p) for w in kel) for p in parcalar)


def _kisaltmali_eslesme(ref: str, arsiv: KanitArsivi, soyad: str) -> KanitKaydi | None:
    """DOI ve cilt/sayfa içermeyen kaynak ("Saygili A, 2018, SIG PROCESS COMMUN"):
    ilk yazar soyadı + yıl + kısaltılmış kaynak adı tek bir yayınımızla eşleşirse o."""
    parca = [p.strip() for p in ref.split(",")]
    if len(parca) < 3 or not soyad or norm(parca[0]).split()[:1] != [norm(soyad)]:
        return None
    yil, kaynak = parca[1], parca[2]
    if not re.fullmatch(r"(19|20)\d{2}", yil) or len(norm(kaynak)) < 4:
        return None
    adaylar = [kay for kay in arsiv.kayitlar if kay.aves_kod and yil in kay.kunye
               and _kisaltma_eslesir(kaynak, kay.kunye.split(ay_baslik(kay.kunye))[-1])]
    return adaylar[0] if len(adaylar) == 1 else None


def ay_baslik(kunye: str) -> str:
    import aves_yardimci as ay
    return ay.baslik_cikar(kunye)


def atif_yapilan_yayinlar(k: WosKaydi, arsiv: KanitArsivi, soyad: str = "") -> list[KanitKaydi]:
    """Kaynakçasında geçen yayınlarımız: DOI; yoksa yıl + cilt + sayfa; o da yoksa
    ilk yazar soyadı + yıl + kısaltılmış kaynak adı (tek yayınla eşleşiyorsa)."""
    bulunan = []
    for ref in k.kaynakca:
        if not doi_bul(ref) and not re.search(r"\bV\d+", ref) and \
                (kay := _kisaltmali_eslesme(ref, arsiv, soyad)) and kay not in bulunan:
            bulunan.append(kay)
    for kay in arsiv.kayitlar:
        if kay in bulunan:
            continue
        if not kay.aves_kod:
            continue
        doiler = {i[4:] for i in kay.kimlikler if i.startswith("doi:")}
        yil, cilt, sayfa = _ref_ipuclari(kay.kunye)
        for ref in k.kaynakca:
            ref_doi = doi_bul(ref)
            if ref_doi and ref_doi in doiler:
                bulunan.append(kay)
                break
            if not ref_doi and yil and cilt and sayfa:
                r = norm(ref)
                if f" {yil} " in f" {r} " and f"v{cilt} " in f"{r} " and f"p{sayfa}" in r.replace(" ", ""):
                    bulunan.append(kay)
                    break
    return bulunan


def oz_atif_mi(k: WosKaydi, soyad: str) -> bool:
    s = norm(soyad)
    return bool(s) and any(norm(a).split()[:1] == [s] for a in k.yazarlar)


def _mevcut_atiflar(kay: KanitKaydi) -> tuple[set[str], list[str]]:
    """Yayının atiflar\\ klasöründeki mevcut atıfların DOI'leri ve ilk sayfa metinleri."""
    doiler, metinler = set(), []
    for alt in kay.atif_klasorleri():
        kj = alt / "kayit.json"
        if kj.exists():
            try:
                if d := json.loads(kj.read_text(encoding="utf-8")).get("doi"):
                    doiler.add(d.lower())
            except (OSError, json.JSONDecodeError):
                pass
        for p in alt.glob("*.pdf"):
            ilk = pdf_metin(p, 1)[:5000]
            if d := doi_bul(ilk):
                doiler.add(d)
            metinler.append(norm(ilk))
    return doiler, metinler


@dataclass
class PlanSatiri:
    wos: WosKaydi
    hedef: KanitKaydi | None
    durum: str               # yeni / zaten var / öz atıf / kapsam dışı / eşleşmedi
    klasor: Path | None = None


def plan_olustur(arsiv: KanitArsivi, kayitlar: list[WosKaydi], soyad: str,
                 kodlar: tuple[str, ...] = ("5.1",)) -> list[PlanSatiri]:
    plan, mevcut = [], {}
    gorulen = set()                                # aynı dışa aktarımda iki kez gelen kayıt
    for w in kayitlar:
        anahtar = w.doi or w.ut or norm(w.baslik)
        if anahtar in gorulen:
            continue
        gorulen.add(anahtar)
        hedefler = atif_yapilan_yayinlar(w, arsiv, soyad)
        if not hedefler:
            plan.append(PlanSatiri(w, None, "eşleşmedi"))
            continue
        for h in hedefler:
            if oz_atif_mi(w, soyad):
                durum = "öz atıf"
            elif w.endeks_kodu not in kodlar:
                durum = "kapsam dışı"
            else:
                if h.aves_kod not in mevcut:
                    mevcut[h.aves_kod] = _mevcut_atiflar(h)
                doiler, metinler = mevcut[h.aves_kod]
                bn = norm(w.baslik)[:80]
                var = (w.doi and w.doi in doiler) or (len(bn) > 30 and any(bn in m for m in metinler))
                durum = "zaten var" if var else "yeni"
            yazar = (w.yazarlar[0].split(",")[0] if w.yazarlar else "Anonim")
            klasor = h.klasor / "atiflar" / (f"WoS_{w.yil or 0}_{slug(yazar, 20)}_"
                                             f"{slug(w.baslik, 45)} ({w.endeks_kisa})")
            plan.append(PlanSatiri(w, h, durum, klasor))
    return plan


# ── Uygulama ───────────────────────────────────────────────────────────────
def _endeks_sayfasi(w: WosKaydi, hedef: KanitKaydi, yol: Path) -> None:
    import fitz
    from .paket import _Yazici
    doc = fitz.open()
    _Yazici(doc).sayfa([
        ("ATIF KANITI – WEB OF SCIENCE KAYDI", 14, True),
        (f"Atıf yapan yayın:\n{w.baslik}", 11, True),
        (f"Yazarlar: {'; '.join(w.yazarlar[:12])}{' …' if len(w.yazarlar) > 12 else ''}", 10, False),
        (f"Dergi: {w.dergi}\nISSN: {w.issn or '-'}   eISSN: {w.eissn or '-'}\n"
         f"Yayın yılı: {w.yil or '-'}   Belge türü: {w.tur or '-'}\n"
         f"DOI: {w.doi or '-'}\nWoS kayıt no (UT): {w.ut or '-'}", 10, False),
        (f"Web of Science Index: {w.endeks or '-'}", 11, True),
        (f"Atıf yapılan eser ({hedef.aves_kod}):\n{hedef.kunye}", 10, False),
        ("Bu sayfa, Web of Science Core Collection 'atıf yapan yayınlar' dışa aktarımından "
         "(Full Record and Cited References) otomatik üretilmiştir. Komisyon ayrıca Master "
         "Journal List ekran görüntüsü isterse aynı klasöre eklenebilir.", 8, False),
    ])
    doc.save(yol)


def uygula(plan: list[PlanSatiri], arsiv: KanitArsivi, indir: bool = True,
           ilerleme=None) -> Counter:
    """'yeni' satırlar için klasör açar; açık erişimli tam metni indirmeyi dener."""
    from .indir import acik_erisim, oturum, pdf_indir
    s = oturum() if indir else None
    sayac = Counter()
    yeniler = [p for p in plan if p.durum == "yeni"]
    for i, p in enumerate(yeniler, 1):
        if ilerleme:
            ilerleme(i, len(yeniler), p)
        w = p.wos
        p.klasor.mkdir(parents=True, exist_ok=True)
        (p.klasor / "kayit.json").write_text(json.dumps({
            "kaynak": KAYNAK, "ut": w.ut, "doi": w.doi, "baslik": w.baslik, "yazarlar": w.yazarlar,
            "dergi": w.dergi, "issn": w.issn, "eissn": w.eissn, "yil": w.yil,
            "endeks": w.endeks, "endeks_kodu": w.endeks_kodu,
            "tarih": w.tarih.isoformat() if w.tarih else None,
            "belge_turu": w.tur, "kitap_bolumu": w.kitap_bolumu,
            "atif_yapilan": p.hedef.aves_kod}, ensure_ascii=False, indent=1), encoding="utf-8")
        _endeks_sayfasi(w, p.hedef, p.klasor / "endeks_bilgisi.pdf")
        indi = False
        if indir and w.doi:
            try:
                for url in acik_erisim(w.doi, s)["pdf"]:
                    ok, _ = pdf_indir(url, p.klasor / "atif_yapan.pdf", s, deneme=2)
                    if ok:
                        indi = True
                        break
            except Exception:  # noqa: BLE001 – ağ hatası klasör oluşturmayı engellemesin
                pass
        if not indi:
            (p.klasor / "INDIRILECEK.txt").write_text(
                "Bu atıf yapan yayının tam metni açık erişimli değil (ya da indirilemedi).\n"
                "Üniversite hesabınızla indirip bu klasöre koyun (dosya adı önemli değil):\n"
                + (f"https://doi.org/{w.doi}\n" if w.doi else f"WoS: {w.ut}\n"), encoding="utf-8")
        sayac["tam metin indirildi" if indi else "tam metin indirilecek"] += 1
    sayac.update(p.durum for p in plan)
    kayitlari_guncelle(arsiv, [p.wos for p in plan])
    indirilecek_listesi(arsiv)
    return sayac


def kayitlari_guncelle(arsiv: KanitArsivi, kayitlar: list[WosKaydi]) -> int:
    """Mevcut WoS atıf klasörlerinin kayit.json'una tarih ve belge türünü işler
    (bu alanlar eklenmeden önce açılmış klasörler için)."""
    ut = {w.ut: w for w in kayitlar if w.ut}
    n = 0
    for kay in arsiv.kayitlar:
        for alt in kay.atif_klasorleri():
            kj = alt / "kayit.json"
            try:
                v = json.loads(kj.read_text(encoding="utf-8")) if kj.exists() else {}
            except (OSError, json.JSONDecodeError):
                continue
            if v.get("kaynak") != KAYNAK or (w := ut.get(v.get("ut"))) is None:
                continue
            yeni = {"tarih": w.tarih.isoformat() if w.tarih else None,
                    "belge_turu": w.tur, "kitap_bolumu": w.kitap_bolumu}
            if any(v.get(k) != d for k, d in yeni.items()):
                v.update(yeni)
                kj.write_text(json.dumps(v, ensure_ascii=False, indent=1), encoding="utf-8")
                n += 1
    return n


def eksikleri_indir(arsiv: KanitArsivi, ilerleme=None) -> Counter:
    """INDIRILECEK.txt bulunan WoS atıf klasörleri için açık erişimli tam metni yeniden dener.

    Bulunan PDF adresi, açık erişim durumu ve sitenin bot koruması kayit.json'a yazılır;
    indirilemeyenler listede tarayıcıdan tek tıkla indirilecek bağlantıyla yer alır."""
    from .indir import oturum, pdf_adaylari, pdf_indir
    bekleyen = [alt for kay in arsiv.kayitlar for alt in kay.atif_klasorleri()
                if (alt / "INDIRILECEK.txt").exists()]
    s, sayac, onbellek = oturum(), Counter(), {}
    for i, alt in enumerate(bekleyen, 1):
        if ilerleme:
            ilerleme(i, len(bekleyen), alt)
        indi = False
        try:
            kj = alt / "kayit.json"
            v = json.loads(kj.read_text(encoding="utf-8"))
            doi = v.get("doi")
            if doi:
                aday = onbellek.get(doi) or pdf_adaylari(doi, s)
                onbellek[doi] = aday
                indi = any(pdf_indir(u, alt / "atif_yapan.pdf", s, deneme=1)[0]
                           for u in aday["pdf"])
                v.update({"acik_erisim": aday["durum"], "pdf_adresi": (aday["pdf"] or [""])[-1],
                          "yayinci_sayfasi": aday["sayfa"], "bot_korumali": aday["korumali"]})
                kj.write_text(json.dumps(v, ensure_ascii=False, indent=1), encoding="utf-8")
        except Exception:  # noqa: BLE001
            indi = False
        sayac["indirildi" if indi else "indirilemedi"] += 1
    indirilecek_listesi(arsiv)                   # indirilenlerin INDIRILECEK.txt'si silinir
    return sayac


def indirilenleri_yerlestir(arsiv: KanitArsivi, klasor: Path, tasi: bool = False) -> Counter:
    """Tarayıcıdan indirilen PDF'leri (örn. İndirilenler klasörü) ilk sayfalarındaki DOI'ye
    ya da başlığa göre bekleyen WoS atıf klasörlerine atif_yapan.pdf olarak koyar.
    Aynı yayın birden çok eserinize atıf yaptıysa her klasöre kopyalanır."""
    import shutil
    bekleyen: dict[str, list[Path]] = {}
    basliklar: dict[str, list[Path]] = {}
    for kay in arsiv.kayitlar:
        for alt in kay.atif_klasorleri():
            if not (alt / "INDIRILECEK.txt").exists():
                continue
            try:
                v = json.loads((alt / "kayit.json").read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError):
                continue
            if v.get("doi"):
                bekleyen.setdefault(v["doi"].lower(), []).append(alt)
            if len(norm(v.get("baslik", ""))) > 25:
                basliklar.setdefault(norm(v["baslik"])[:80], []).append(alt)
    sayac = Counter()
    for p in sorted(Path(klasor).glob("*.pdf")):
        metin = pdf_metin(p, 2)[:8000]
        hedefler = bekleyen.get(doi_bul(metin)) or next(
            (v for b, v in basliklar.items() if b in norm(metin)), None)
        if not hedefler:
            sayac["eşleşmedi"] += 1
            continue
        for alt in hedefler:
            if not (alt / "atif_yapan.pdf").exists():
                shutil.copy2(p, alt / "atif_yapan.pdf")
                sayac["yerleştirildi"] += 1
        if tasi:
            p.unlink()
    indirilecek_listesi(arsiv)
    return sayac


def indirilecek_listesi(arsiv: KanitArsivi) -> Path | None:
    """Tam metni henüz konmamış WoS atıflarının listesi (kanıt kökünde Excel)."""
    satirlar = []
    for kay in arsiv.kayitlar:
        for alt in kay.atif_klasorleri():
            if (alt / "INDIRILECEK.txt").exists():
                pdfler = [p for p in alt.glob("*.pdf") if p.name != "endeks_bilgisi.pdf"]
                if pdfler:                         # kullanıcı tam metni koymuş
                    (alt / "INDIRILECEK.txt").unlink()
                    continue
                try:
                    v = json.loads((alt / "kayit.json").read_text(encoding="utf-8"))
                except (OSError, json.JSONDecodeError):
                    v = {}
                oa = v.get("acik_erisim")
                durum = ("Açık erişim – tarayıcıdan indirin" if oa and oa != "closed"
                         else "Ücretli – kurum erişimi gerekir" if oa == "closed" else "")
                satirlar.append([kay.aves_kod, v.get("baslik", alt.name), v.get("dergi", ""),
                                 v.get("yil", ""), durum,
                                 v.get("pdf_adresi") or "",
                                 f"https://doi.org/{v['doi']}" if v.get("doi") else "",
                                 str(alt)])
    yol = arsiv.kok / INDIRILECEK_LISTESI
    if not satirlar:
        try:
            yol.unlink(missing_ok=True)
        except PermissionError:          # Excel'de açık: bir sonraki güncellemede silinir
            pass
        return None
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "İndirilecek atıflar"
    ws.append(["Atıf yapılan", "Atıf yapan yayın", "Dergi", "Yıl", "Erişim", "PDF bağlantısı",
               "DOI bağlantısı", "Klasör"])
    # açık erişimliler üstte: tarayıcıda tıklayıp indirmek yeterli
    satirlar.sort(key=lambda r: (not r[4].startswith("Açık"), r[0]))
    for r in satirlar:
        ws.append(r)
        for sutun in (6, 7):
            h = ws.cell(ws.max_row, sutun)
            if h.value:
                h.hyperlink = h.value
                h.style = "Hyperlink"
    for col, w in zip("ABCDEFGH", (10, 60, 30, 6, 30, 45, 35, 70)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    try:
        wb.save(yol)
    except PermissionError:
        # Liste Excel'de açıkken Windows dosyayı kilitler: güncel liste yanına yazılır
        yol = yol.with_name(yol.stem + " (güncel).xlsx")
        try:
            wb.save(yol)
        except PermissionError:
            return None
    return yol
