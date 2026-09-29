"""Kanıt arşivinin özeti: kunye.txt kontrol listelerini ve Excel özetini günceller."""

from __future__ import annotations

import re
from pathlib import Path

from .kayit import KUNYE_DOSYASI, KanitArsivi
from .kontrol import KANIT_TURLERI, durum

OZET_DOSYASI = "Faaliyet_Kanit_Ozeti.xlsx"


def kunye_guncelle(k, d) -> None:
    """kunye.txt'nin ilk bölümünü korur; kontrol listesi ve arşiv bölümünü yeniden yazar."""
    yol = k.klasor / KUNYE_DOSYASI
    if not yol.exists():
        return
    bas = yol.read_text(encoding="utf-8").split("\n\nBaşvuru dosyası için")[0].rstrip()
    tam = "VAR" if k.tam_metin else "YOK"
    bas = re.sub(r"^Tam metin    : .*$",
                 f"Tam metin    : {tam}" + (f" ({k.tam_metin.name})" if k.tam_metin else ""),
                 bas, flags=re.M)
    liste = []
    for tur in list(d.bulunan) + d.eksik:
        if tur == "tam_metin":
            continue
        dosyalar = d.bulunan.get(tur, [])
        liste.append(f"  [{'x' if dosyalar else ' '}] {KANIT_TURLERI[tur][0]}"
                     + (f"   ← {', '.join(dosyalar[:3])}" if dosyalar else ""))
    atiflar = k.atif_klasorleri()
    ek = ["", "── Arşivden eklenenler ─────────────────────────────",
          f"Arşiv dosyası : {len(k.arsiv_dosyalari())}  (arsiv\\ alt klasörü)",
          f"Atıf kanıtı   : {len(atiflar)} atıf  (atiflar\\ alt klasörü)"]
    ek += [f"   • {a.name}" for a in atiflar[:60]]
    yol.write_text(bas + "\n\nBaşvuru dosyası için ayrıca gerekli kanıtlar "
                   "(otomatik işaretlendi – elle doğrulayın):\n" + "\n".join(liste)
                   + "\n" + "\n".join(ek) + "\n", encoding="utf-8")


def ozet_yaz(arsiv: KanitArsivi, excel: bool = True) -> list[list]:
    satirlar = []
    for k in sorted(arsiv.kayitlar, key=lambda k: k.klasor.name):
        d = durum(k)
        kunye_guncelle(k, d)
        satirlar.append([k.aves_kod, k.kunye, k.tur, k.doi,
                         k.tam_metin.name if k.tam_metin else "YOK",
                         len(k.arsiv_dosyalari()), len(k.atif_klasorleri()),
                         "; ".join(d.eksik_aciklamalari()) or "—", k.klasor.name])
    if excel:
        _excel_yaz(arsiv, satirlar)
    return satirlar


def _excel_yaz(arsiv: KanitArsivi, satirlar):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Yayınlar"
    ws.append(["Kod", "Künye (AVES)", "Endeks (AVES)", "DOI", "Tam Metin", "Arşiv Dosyası",
               "Atıf Kanıtı", "Eksik Kanıtlar", "Klasör"])
    for r in satirlar:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2E5DA3")
    for row in ws.iter_rows(min_row=2):
        row[4].fill = PatternFill("solid", fgColor="FFEBEE" if row[4].value == "YOK" else "E8F5E9")
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate([7, 70, 12, 26, 20, 9, 9, 50, 40], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet("Diğer Klasörler")
    ws2.append(["Klasör", "Dosya", "MB"])
    for d in arsiv.diger_klasorler:
        dosyalar = [p for p in d.rglob("*") if p.is_file()]
        ws2.append([d.name, len(dosyalar), round(sum(p.stat().st_size for p in dosyalar) / 1e6, 1)])
    ws2.column_dimensions["A"].width = 45
    wb.save(arsiv.kok / OZET_DOSYASI)
